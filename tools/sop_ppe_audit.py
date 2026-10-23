# -*- coding: utf-8 -*-
"""작업별 개인보호구가 맞게 들어갔는지 훑어본다. (읽기만 — 아무것도 고치지 않는다)

어떻게 보나
    작업명 · 조치사항 · 각부 설명에 적힌 **실제 글**에서 위험요인을 찾고,
    그 위험요인에 필요한 보호구와 지금 적힌 보호구를 견준다.

    · 빠짐  : 위험요인이 있는데 보호구가 없다
    · 과함  : 보호구는 있는데 그럴 위험요인이 안 보인다
    · 주의  : 그 작업에서 오히려 위험한 보호구 (회전체 작업의 장갑 등)

    판정은 제안이다. 최종 판단은 사람이 한다.

쓰는 법
    python sop_ppe_audit.py                  화면으로 요약
    python sop_ppe_audit.py --all            작업별 상세까지
    python sop_ppe_audit.py --xlsx 결과.xlsx  검토용 엑셀로
"""
import argparse
import io
import json
import os
import re
import subprocess
import sys

SITE = (r"C:\Users\user\Desktop\Bio Innovation Hub\01 Projects"
        r"\[안전] 표준운영절차 SOP\웹 배포\안전SOP_웹사이트")

# ── 위험요인 : (이름, 찾을 낱말, 그래서 필요한 보호구)
위험 = [
    ("고소작업", ["고소", "사다리", "비계", "고소작업대", "SKY", "고소작업차",
                "옥상", "상층부", "지붕", "추락", "달비계", "리프트"],
     ["안전대"]),
    ("활선·충전부", ["활선", "충전부", "검전", "정전작업", "수전", "배전반",
                  "MCC", "분전반", "차단기", "VCB", "ACB", "변압기", "큐비클",
                  "비상발전기", "UPS", "배터리", "콘센트", "결선", "단자대",
                  "수변전", "변전", "특고압", "고압", "저압", "전기실",
                  "발전기", "아크", "용접기",
                  "전선 교체", "등기구 교체", "안정기"],
     ["절연장갑"]),
    ("질식·산소결핍", ["가스 소화", "가스계", "이산화탄소", "청정약제", "할론",
                   "CO2", "질식", "산소결핍", "약제 방출"],
     ["공기호흡기"]),
    ("물·습식", ["세척", "수조", "저수조", "배수", "급수", "물청소", "세정",
               "집수정", "오수", "펌프실", "소화수"],
     ["고무장갑"]),
    ("화기·용접", ["용접", "화기", "그라인더", "연삭", "절단", "산소", "토치"],
     ["차광보안면", "가죽장갑"]),
    ("밀폐공간", ["밀폐", "맨홀", "피트", "탱크 내부", "저수조 내부", "덕트 내부"],
     ["송기마스크"]),
    ("소음", ["송풍기", "냉각탑", "발전기", "펌프", "공조기", "팬룸", "블로워"],
     ["귀마개"]),
    ("분진", ["필터", "분진", "먼지", "연삭", "그라인딩", "석면", "집진"],
     ["방진마스크"]),
    ("화학물질", ["약품", "약액", "세제", "왁스", "수처리", "청관제", "살균",
                "차아염소", "유해화학", "유해물질", "흡입", "유기용제"],
     ["방독마스크", "보안경"]),
    ("고온·화상", ["화상", "고온", "스팀", "증기", "열수", "용융", "급탕"],
     ["내열장갑"]),
    ("회전체", ["벨트", "풀리", "그리스", "커플링", "회전체", "임펠러", "축"],
     []),
]

# 이름이 달라도 같은 것으로 본다
같은것 = {
    "안전대": ["안전대", "안전벨트", "전신안전대", "추락방지대"],
    "절연장갑": ["절연장갑"],
    "고무장갑": ["고무장갑", "내화학 보호장갑", "방수장갑"],
    "보호장갑": ["보호장갑", "안전장갑", "작업장갑", "면장갑"],
    "차광보안면": ["차광보안면", "용접면", "용접용 보안면"],
    "안면보호구": ["안면보호구", "보호면", "페이스실드"],
    "가죽장갑": ["가죽장갑", "용접장갑"],
    "송기마스크": ["송기마스크", "에어라인"],
    "공기호흡기": ["공기호흡기", "SCBA"],
    "귀마개": ["귀마개", "귀덮개", "방음보호구"],
    "방진마스크": ["방진마스크"],
    "방독마스크": ["방독마스크"],
    "보안경": ["보안경", "고글"],
    "안전모": ["안전모"],
    "안전화": ["안전화", "논슬립 안전화", "절연화", "장화"],
}


def 묶음(이름):
    """'귀마개 · 보안경' 처럼 붙어 있는 것도 쪼개서 표준 이름으로."""
    out = set()
    for 조각 in re.split(r"[·,/]|\s·\s", 이름):
        조각 = 조각.strip()
        if not 조각:
            continue
        for 표준, 말들 in 같은것.items():
            if any(w in 조각 for w in 말들):
                out.add(표준)
    return out


# 이런 말이 같은 문장에 있으면 '하는 작업' 이 아니라 '하지 말라는 말' 이다.
금지말 = ["금지", "엄금", "금함", "말 것", "하지 마", "불가", "없도록",
        "방지", "주의", "위험", "확인", "제거", "차단", "소화기", "비치"]


def 작업글(t):
    """그 작업에 적힌 글을 문장 목록으로. 작업명은 빼고 돌려준다."""
    조각 = [t.get("place", "")]
    for sec in t.get("guide") or []:
        조각.append(sec.get("title", ""))
        for it in sec.get("items") or []:
            조각.append(it.get("title", ""))
            조각 += it.get("bullets") or []
    p = t.get("parts") or {}
    for it in p.get("items") or []:
        조각 += [it.get("name", ""), it.get("desc", "")]
    for k in ("pre", "use"):
        조각 += p.get(k) or []
    # ban(금지 사항)은 통째로 뺀다 — 하지 말라는 말이다
    return [x for x in 조각 if x]


def 보다(T):
    결과 = []
    for tid, t in T.items():
        문장들 = 작업글(t)
        이름글 = t.get("name", "")
        있는것 = set()
        원래 = [x["name"] for x in (t.get("ppe") or [])]
        for n in 원래:
            있는것 |= 묶음(n)

        # 작업명에 나오면 확실한 위험요인. 본문은 금지 문장을 뺀 뒤
        # 서로 다른 낱말이 2개 이상 맞을 때만 인정한다.
        찾은위험, 필요, 약한위험 = [], set(), []
        for 이름, 낱말, 보호구 in 위험:
            이름맞음 = [w for w in 낱말 if w in 이름글]
            본문맞음 = []
            for 문장 in 문장들:
                if any(g in 문장 for g in 금지말):
                    continue
                본문맞음 += [w for w in 낱말 if w in 문장]
            본문맞음 = sorted(set(본문맞음))
            if 이름맞음 or len(본문맞음) >= 2:
                찾은위험.append((이름, (이름맞음 or 본문맞음)[:3]))
                필요 |= set(보호구)
            elif 본문맞음:
                약한위험.append((이름, 본문맞음[:2]))

        위험이름 = {x[0] for x in 찾은위험}
        애매 = 위험이름 | {x[0] for x in 약한위험}

        빠짐 = sorted(필요 - 있는것)

        과함 = []
        if "절연장갑" in 있는것 and "활선·충전부" not in 애매:
            과함.append("절연장갑 → 보호장갑 (전기 위험요인 없음)")
        if "고무장갑" in 있는것 and not ({"물·습식", "화학물질"} & 애매):
            과함.append("고무장갑 → 보호장갑 (물·약품 위험요인 없음)")
        if "송기마스크" in 있는것 and "밀폐공간" not in 애매:
            과함.append("송기마스크 (밀폐공간 아님)")
        if "공기호흡기" in 있는것 and not ({"밀폐공간", "질식·산소결핍"} & 애매):
            과함.append("공기호흡기 (밀폐·질식 위험요인 없음)")
        if "차광보안면" in 있는것 and "화기·용접" not in 애매:
            과함.append("차광보안면 (화기·용접 아님)")

        주의 = []
        if "회전체" in 위험이름 and ("보호장갑" in 있는것 or "고무장갑" in 있는것):
            주의.append("회전체 작업 — 장갑 말림 위험. 착용 금지 구간 명시 필요")
        if "고소작업" in 위험이름 and "안전모" in 있는것:
            주의.append("고소작업 — 안전모 턱끈 체결 명시 필요")

        결과.append({
            "id": tid, "name": t.get("name", ""), "place": t.get("place", ""),
            "ppe": 원래, "위험": 찾은위험,
            "빠짐": 빠짐, "과함": 과함, "주의": 주의,
        })
    return 결과


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--all", action="store_true", help="작업별 상세까지")
    ap.add_argument("--xlsx", help="검토용 엑셀로 저장")
    a = ap.parse_args()

    p = os.path.join(SITE, "assets", "data.js").replace("\\", "/")
    js = ("global.window={};require(%s);"
          "process.stdout.write(JSON.stringify(window.TASKS))" % json.dumps(p))
    T = json.loads(subprocess.run(["node", "-e", js], capture_output=True,
                                  check=True).stdout.decode("utf-8"))
    R = 보다(T)

    빠진것 = [r for r in R if r["빠짐"]]
    과한것 = [r for r in R if r["과함"]]
    주의것 = [r for r in R if r["주의"]]
    깨끗 = [r for r in R if not (r["빠짐"] or r["과함"])]

    print("작업 %d건을 봤습니다.\n" % len(R))
    print("  빠진 보호구가 있는 작업 : %3d건" % len(빠진것))
    print("  과하게 들어간 작업      : %3d건" % len(과한것))
    print("  착용 주의 안내가 필요    : %3d건" % len(주의것))
    print("  그대로 둬도 되는 작업    : %3d건" % len(깨끗))

    from collections import Counter
    c = Counter(x for r in 빠진것 for x in r["빠짐"])
    print("\n■ 가장 많이 빠진 보호구")
    for 이름, n in c.most_common():
        print("   %-10s %3d건" % (이름, n))

    c2 = Counter(x.split(" ")[0] for r in 과한것 for x in r["과함"])
    print("\n■ 가장 많이 과하게 들어간 보호구")
    for 이름, n in c2.most_common():
        print("   %-10s %3d건" % (이름, n))

    if a.all:
        print("\n" + "=" * 70)
        for r in sorted(R, key=lambda x: (not x["빠짐"], not x["과함"], x["id"])):
            if not (r["빠짐"] or r["과함"] or r["주의"]):
                continue
            print("\n[%s] %s  (%s)" % (r["id"], r["name"], r["place"]))
            print("   지금 보호구 : %s" % ", ".join(r["ppe"]))
            print("   찾은 위험   : %s"
                  % ", ".join("%s(%s)" % (n, "·".join(w)) for n, w in r["위험"]))
            for x in r["빠짐"]:
                print("   [빠짐] %s" % x)
            for x in r["과함"]:
                print("   [과함] %s" % x)
            for x in r["주의"]:
                print("   [주의] %s" % x)

    if a.xlsx:
        쓰기(R, a.xlsx)
        print("\n엑셀 저장 : %s" % a.xlsx)


def 쓰기(R, path):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

    F = "맑은 고딕"
    NAVY, YEL, RED, SOFT = "1E3A6E", "FFF6D5", "B3261E", "F6F8FB"
    thin = Side(style="thin", color="B9C2CF")
    med = Side(style="medium", color=NAVY)
    BOX = Border(left=thin, right=thin, top=thin, bottom=thin)
    NL = chr(10)

    wb = Workbook()
    ws = wb.active
    ws.title = "보호구 점검"
    ws.sheet_view.showGridLines = False

    ws.merge_cells("B2:J2")
    ws["B2"] = "작업별 개인보호구 점검 — 빠진 것 · 과한 것"
    ws["B2"].font = Font(name=F, size=16, bold=True, color=NAVY)
    ws.row_dimensions[2].height = 30
    ws.merge_cells("B3:J3")
    ws["B3"] = ("작업 %d건 · 작업명과 조치사항 글에서 위험요인을 찾아 대조한 결과입니다.　"
                "판정은 제안이며 최종 결정은 '결정' 칸에 적어 주세요." % len(R))
    ws["B3"].font = Font(name=F, size=10.5, color="5A6B82")

    HEAD = ["번호", "작업명", "장소", "지금 보호구", "찾은 위험요인",
            "빠진 것", "과한 것", "착용 주의", "결정", "비고"]
    r = 5
    for i, h in enumerate(HEAD):
        c = ws.cell(row=r, column=2 + i, value=h)
        c.font = Font(name=F, size=10.5, bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor=NAVY)
        c.alignment = Alignment(horizontal="center", vertical="center",
                                wrap_text=True)
        c.border = Border(left=thin, right=thin, top=med, bottom=med)
    ws.row_dimensions[r].height = 28

    r = 6
    for x in sorted(R, key=lambda v: (not v["빠짐"], not v["과함"], v["id"])):
        문제 = bool(x["빠짐"] or x["과함"])
        vals = [x["id"], x["name"], x["place"], ", ".join(x["ppe"]),
                NL.join("%s (%s)" % (n, "·".join(w)) for n, w in x["위험"]),
                NL.join(x["빠짐"]), NL.join(x["과함"]), NL.join(x["주의"]),
                None, None]
        for i, v in enumerate(vals):
            c = ws.cell(row=r, column=2 + i, value=v)
            색 = RED if (i in (5, 6) and v) else "1B2430"
            c.font = Font(name=F, size=9.5, bold=(i == 0), color=색)
            c.border = BOX
            c.fill = PatternFill("solid",
                                 fgColor=YEL if i in (8, 9)
                                 else (SOFT if 문제 else "FFFFFF"))
            c.alignment = Alignment(
                horizontal="center" if i in (0, 2) else "left",
                vertical="center", wrap_text=True, indent=0 if i == 0 else 1)
        줄 = max([str(v).count(NL) + 1 for v in vals if v] or [1])
        ws.row_dimensions[r].height = max(26, 13 * 줄 + 8)
        r += 1

    for col, w in (("A", 2.5), ("B", 8), ("C", 30), ("D", 16), ("E", 28),
                   ("F", 26), ("G", 16), ("H", 30), ("I", 30), ("J", 16),
                   ("K", 20)):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "B6"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.print_title_rows = "5:5"
    wb.save(path)


if __name__ == "__main__":
    main()
