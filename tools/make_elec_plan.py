from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH as PA
from docx.enum.table import WD_ALIGN_VERTICAL as VA
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

FONT = '맑은 고딕'
DARK  = '1F3864'
LBLUE = 'BDD7EE'
MBLUE = 'D9E1F2'
WHITE = 'FFFFFF'
PW    = 17.4

def ef(n): return OxmlElement(n)

def set_bg(cell, hex_color):
    shd = ef('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), hex_color)
    cell._tc.get_or_add_tcPr().append(shd)

def set_borders(cell, sz='4'):
    tcPr = cell._tc.get_or_add_tcPr()
    b = ef('w:tcBorders')
    for edge in ('top','left','bottom','right'):
        e = ef(f'w:{edge}')
        e.set(qn('w:val'), 'single')
        e.set(qn('w:sz'), sz)
        e.set(qn('w:space'), '0')
        e.set(qn('w:color'), '000000')
        b.append(e)
    tcPr.append(b)

def all_b(tbl, sz='4'):
    for row in tbl.rows:
        for cell in row.cells:
            set_borders(cell, sz)

def set_fnt(run, sz, bold, fg):
    run.font.name = FONT
    run.font.size = Pt(sz)
    run.font.bold = bold
    if fg:
        run.font.color.rgb = RGBColor.from_string(fg)
    rPr = run._r.get_or_add_rPr()
    f = ef('w:rFonts')
    for attr in ('w:ascii', 'w:eastAsia', 'w:hAnsi'):
        f.set(qn(attr), FONT)
    rPr.insert(0, f)

def cw(cell, txt, sz=9, bold=False, align=PA.LEFT, va=VA.CENTER, fg=None):
    cell.vertical_alignment = va
    p = cell.paragraphs[0]
    p.alignment = align
    p.paragraph_format.space_before = Pt(1)
    p.paragraph_format.space_after  = Pt(1)
    r = p.add_run(txt)
    set_fnt(r, sz, bold, fg)

def ca(cell, txt, sz=9, bold=False, align=PA.LEFT, fg=None):
    p = cell.add_paragraph()
    p.alignment = align
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after  = Pt(0)
    r = p.add_run(txt)
    set_fnt(r, sz, bold, fg)

def dark_hdr(cell, txt):
    set_bg(cell, DARK)
    cw(cell, txt, sz=10, bold=True, fg=WHITE)

def col_hdr(cell, txt):
    set_bg(cell, LBLUE)
    cw(cell, txt, sz=9, align=PA.CENTER)

def lbl(cell, txt):
    set_bg(cell, MBLUE)
    cw(cell, txt, sz=9, align=PA.CENTER)

def sw(tbl, widths):
    for row in tbl.rows:
        for i, cell in enumerate(row.cells):
            if i < len(widths):
                cell.width = Cm(widths[i])

def mk(doc, rows, cols, widths, sz='4'):
    t = doc.add_table(rows=rows, cols=cols)
    t.style = 'Table Grid'
    all_b(t, sz)
    sw(t, widths)
    return t

def tight(doc):
    for p in doc.paragraphs:
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after  = Pt(0)

# ── Document setup ──────────────────────────────────────────
doc = Document()
sec = doc.sections[0]
sec.page_width    = Cm(21)
sec.page_height   = Cm(29.7)
sec.left_margin   = Cm(1.8)
sec.right_margin  = Cm(1.8)
sec.top_margin    = Cm(1.5)
sec.bottom_margin = Cm(1.5)
ns = doc.styles['Normal']
ns.font.name = FONT
ns.font.size = Pt(9)
ns.paragraph_format.space_before = Pt(0)
ns.paragraph_format.space_after  = Pt(0)

# ═══════════════════════════════════════════════════
# PAGE 1
# ═══════════════════════════════════════════════════

# 상단 헤더
t = mk(doc, 3, 4, [10.5, 2.0, 2.0, 2.9], sz='6')
mc = t.cell(0,0).merge(t.cell(2,0))
cw(mc, '전기작업계획서', sz=16, bold=True, align=PA.CENTER)
for j, h in enumerate(['담당','검토','승인'], 1):
    cw(t.cell(0,j), h, sz=9, align=PA.CENTER)
t.rows[1].height = Cm(1.2)
cw(t.cell(2,1), '작성일자', sz=9, align=PA.CENTER)
t.cell(2,2).merge(t.cell(2,3))
cw(t.cell(2,2), '2026. 1. 1.', sz=9, align=PA.CENTER)
tight(doc)

# 섹션1 헤더
t1h = mk(doc, 1, 1, [PW], sz='6')
dark_hdr(t1h.rows[0].cells[0], '1. 작업(공사) 개요')
tight(doc)

# 섹션1 내용
s1_rows = [
    '■ 공 사 명 :    건물 내 전기 작업',
    '■ 공사기간 :    2026. 1. 1.  ~  2026. 12. 31.',
    '■ 발 주 처 :    한독제넥신프로젠 연구소              ○ 담당자 :           (연락처)',
    '■ 공사업체 :    삼구아이앤씨                         ○ 작업책임자 : 김희진  (연락처) 010-7132-5105',
    '■ 작업목적 및 내용 :    건물 전기 관련 작업',
    '■ 공 사 범 위 :    조명 교체 작업, 비상발전기 가동, 콘센트 증설, 분전함 점검',
    '■ 작 업 자 :    (  5  명)',
]
t1c = mk(doc, len(s1_rows), 1, [PW])
for i, txt in enumerate(s1_rows):
    cw(t1c.rows[i].cells[0], txt)
tight(doc)

# 작업자 테이블
tw = mk(doc, 4, 6, [1.5, 3.5, 3.5, 1.5, 3.5, 3.9])
for j, h in enumerate(['연번','성명','비고(자격)','연번','성명','비고(자격)']):
    col_hdr(tw.rows[0].cells[j], h)
workers = [
    ['1','서동혁','전기기사','4','임채혁','-'],
    ['2','박재범','전기기능사','5','이성호','-'],
    ['3','김주환','-','6','',''],
]
for ri, row in enumerate(workers):
    for ci, v in enumerate(row):
        cw(tw.rows[ri+1].cells[ci], v, align=PA.CENTER)
tight(doc)

# 섹션2 헤더
t2h = mk(doc, 1, 1, [PW], sz='6')
dark_hdr(t2h.rows[0].cells[0], '2. 안전보건교육 및 보호구 지급 계획 등')
tight(doc)

# 교육계획 소제목
t2s = mk(doc, 1, 1, [PW])
cw(t2s.rows[0].cells[0], '■ 안전보건교육 계획')
tight(doc)

# 교육 테이블
te = mk(doc, 3, 4, [2.2, 4.8, 2.2, 8.2])
lbl(te.rows[0].cells[0], '교육일자')
cw(te.rows[0].cells[1], '2026. 2. 20.', align=PA.CENTER)
lbl(te.rows[0].cells[2], '교육장소')
cw(te.rows[0].cells[3], '방재실', align=PA.CENTER)
lbl(te.rows[1].cells[0], '교육자')
cw(te.rows[1].cells[1], '서동혁', align=PA.CENTER)
lbl(te.rows[1].cells[2], '인원')
cw(te.rows[1].cells[3], '5', align=PA.CENTER)
lbl(te.rows[2].cells[0], '교육내용')
mc_edu = te.cell(2,1).merge(te.cell(2,3))
cw(mc_edu, '1. 전기작업이란?   2. 전기 작업 시 안전보건 주의사항')
ca(mc_edu, '3. 재해사례         4. 관련 법령 등')
tight(doc)

# 보호구 소제목
t2s2 = mk(doc, 1, 1, [PW])
cw(t2s2.rows[0].cells[0], '■ 보호구 지급 및 사용 장비')
tight(doc)

# 보호구/장비 테이블
teq = mk(doc, 7, 3, [2.0, 3.2, 12.2])
mc_jp = teq.cell(0,0).merge(teq.cell(1,0))
lbl(mc_jp, '지급\n보호구')
mc0 = teq.cell(0,1).merge(teq.cell(0,2))
cw(mc0, '■절연안전모(ABE)   ■안전화/안전모   ■절연장화   ■안전대   ■보안경')
mc1 = teq.cell(1,1).merge(teq.cell(1,2))
cw(mc1, '■절연장갑   ■방진마스크   □기타(                              )')
mc_n = teq.cell(2,0).merge(teq.cell(2,2))
cw(mc_n, '* 보호구 지급대장 관리 및 사용전 점검하여 이상시 교체', sz=8)
mc_eq = teq.cell(3,0).merge(teq.cell(6,0))
lbl(mc_eq, '사용\n장비')
for ri, (sub, cont) in enumerate([
    ('측정장비', '멀티테스터기, 검진기'),
    ('활선기구/장비', ''),
    ('방호구', ''),
    ('기타', ''),
]):
    lbl(teq.rows[3+ri].cells[1], sub)
    cw(teq.rows[3+ri].cells[2], cont)
tight(doc)

# ═══════════════════════════════════════════════════
# PAGES 2-5
# ═══════════════════════════════════════════════════
pages = [
    {
        'work': '조명\n교체\n작업',
        'rows': [
            ('전원차단', '전원 미차단 시 감전',
             '1. 해당 구역 차단기 내림 및\n"작업중" 표지판 부착\n2. 절연 장갑 및 절연 공구 사용', ''),
            ('노후 등기구\n해체', '사다리 작업 중 추락',
             '1. 2인 1조 작업 및 안전모 착용\n2. 사다리 최상단 작업 금지 및\n아웃트리거 확인', ''),
            ('신규 등기구\n설치 및 결선', '등기구 낙하로 인한 맞음',
             '2인 1조 및 올바른 자세로 작업', ''),
            ('점등 확인', '', '', ''),
        ],
        'si': ['■','■','□','□'], 'sa': ['□','■','■','□'],
    },
    {
        'work': '비상\n발전기\n가동',
        'rows': [
            ('가동 전 육안\n점검(유량 등)', '회전체(팬, 벨트)에 끼임',
             '가동 전 주변 가연물 제거 및\n환기 시설 가동', ''),
            ('수동/자동\n가동 시험', '엔진 가동 시 화상 및 소음',
             '1. 회전 부위 덮개 설치 확인 및\n접근 금지\n2. 청력 보호구(귀덮개 등) 착용', ''),
            ('전압 및\n주파수 확인', '배기가스에 의한 질식/중독',
             '가동 중인 엔진 몸체 접촉 금지', ''),
            ('정지 및 복구', '', '', ''),
        ],
        'si': ['■','■','□','□'], 'sa': ['□','■','■','□'],
    },
    {
        'work': '콘센트\n증설\n작업',
        'rows': [
            ('회로 확인 및\n전원 차단', '',
             '분전반 내 해당 차단기 개방 및\n잠금장치(LOTO) 실시', ''),
            ('배관 설치 및\n입선(배선) 작업', '활선 작업 시 감전',
             '1. 검진기 사용하여 무전입 상태\n확인 후 작업\n2. 절연장갑 착용 후 작업', ''),
            ('콘센트 기구\n결선 및 고정',
             '1. 부적절한 결선으로 인한\n아크/화재\n2. 절연장갑 미착용으로\n인한 감전',
             '절연장갑 착용 후 작업', ''),
            ('전압 측정 및\n통전 테스트', '작업 중 전원 투입 시\n작업자 감전 위험',
             '작업 마무리 확인 후 전원 투입', ''),
        ],
        'si': ['□','□','□','□'], 'sa': ['□','□','□','□'],
    },
    {
        'work': '분전함\n점검',
        'rows': [
            ('입함 변형 및\n부식 확인', '충전부 노출에 의한 감전',
             '1. 문 개방 전 검전기로 외함\n누전 여부 확인\n2. 절연용 보호구 착용', ''),
            ('내부 배선 및\n단자 체결 확인', '단자 이완으로 인한\n아크/화재',
             '내부 배선 점검 일지 작성', ''),
            ('절연 저항 및\n부하 측정', '', '', ''),
            ('내부 청소 및\n복구', '이물질 제거 중 단락',
             '비전도성 청소도구 사용', ''),
        ],
        'si': ['■','■','□','□'], 'sa': ['□','■','■','□'],
    },
]

NOTE_LINES = [
    '※ 위 내용에는 다음 사항이 포함되어야 함',
    '① 작업의 목적 및 내용, ② 작업자의 자격 및 적정인원, ③ 작업 범위, 작업책임자 임명, 전격·아크 섬광·아크 폭발 등 전기 위험 요인 파악, 접근 한계거리, 활선접근 경보장치 휴대 등 작업시작 전에 필요한 사항,',
    '④ 전로차단에 관한 작업계획 및 전원(電源) 재투입 절차 등 작업 상황에 필요한 안전 작업 요령, ⑤ 절연용 보호구 및 방호구, 활선작업기구·장치 등의 준비·점검·착용·사용 등에 관한 사항,',
    '⑥ 점검·시운전을 위한 일시 운전, 작업 중단 등에 관한 사항, ⑦ 교대 근무 시 근무 인계(引繼)에 관한 사항, ⑧ 전기작업 장소에 대한 관계 근로자가 아닌 사람의 출입금지에 관한 사항,',
    '⑨ 전기안전작업계획서를 해당 근로자에게 교육할 수 있는 방법과 작성된 전기안전작업계획서의 평가·관리계획, ⑩ 전기 도면, 기기 세부 사항 등 작업과 관련되는 자료 등',
]

for idx, pg in enumerate(pages):
    doc.add_page_break()
    tight(doc)

    # 연속 헤더
    th = mk(doc, 3, 3, [10.0, 3.2, 4.2], sz='6')
    mc_ti = th.cell(0,0).merge(th.cell(2,0))
    cw(mc_ti, '전기작업계획서', sz=14, bold=True, align=PA.CENTER)
    for ri, (lab, val) in enumerate([
        ('작업공정(부서)', '전기파트'),
        ('작 성 자', '서동혁'),
        ('작성일자', '2026. 1. 1.'),
    ]):
        lbl(th.rows[ri].cells[1], lab)
        cw(th.rows[ri].cells[2], val, align=PA.CENTER)
    tight(doc)

    # 섹션3 헤더
    t3h = mk(doc, 1, 1, [PW], sz='6')
    dark_hdr(t3h.rows[0].cells[0], '3. 작업순서 및 안전작업 방법')
    tight(doc)

    # 섹션3 내용
    nr = len(pg['rows'])
    t3 = mk(doc, 1+nr, 5, [2.0, 3.3, 3.0, 7.2, 1.9])
    for j, h in enumerate(['작업명','작업순서 및\n작업내용','위험요소','안전작업방법','비고']):
        col_hdr(t3.rows[0].cells[j], h)
    mc_wt = t3.cell(1,0).merge(t3.cell(nr,0))
    lbl(mc_wt, pg['work'])
    for ri, (step, hazard, method, note) in enumerate(pg['rows']):
        cw(t3.rows[1+ri].cells[1], step, align=PA.CENTER)
        cw(t3.rows[1+ri].cells[2], hazard)
        cw(t3.rows[1+ri].cells[3], method)
        cw(t3.rows[1+ri].cells[4], note, align=PA.CENTER)
    tight(doc)

    # 섹션4
    t4h = mk(doc, 1, 1, [PW], sz='6')
    dark_hdr(t4h.rows[0].cells[0], '4. 현장 통제 및 안전조치 계획')
    tight(doc)
    t4c = mk(doc, 1, 1, [PW])
    c4 = t4c.rows[0].cells[0]
    cw(c4, '1. 전기 작업 시 전원 투입기구에 LOTO 조치 실시')
    ca(c4, '2. 전원 투입구에 인원 배치하여 오조작으로 인한 감전사고 예방')
    ca(c4, '3. 전기 작업중을 알리는 표지판 설치')
    tight(doc)

    # 섹션5
    t5h = mk(doc, 1, 1, [PW], sz='6')
    dark_hdr(t5h.rows[0].cells[0], '5. 첨부 자료')
    tight(doc)
    si, sa = pg['si'], pg['sa']
    t5 = mk(doc, 4, 2, [1.8, 15.6])
    mc_sb = t5.cell(0,0).merge(t5.cell(1,0))
    lbl(mc_sb, '설비')
    cw(t5.rows[0].cells[1],
       f'{si[0]}전기단선도/결선도   {si[1]} 기기 사양서   {si[2]}공사계획표(정전작업 계획 등)')
    ca(t5.rows[0].cells[1], f'{si[3]}기타 (                    )')
    mc_sa = t5.cell(2,0).merge(t5.cell(3,0))
    lbl(mc_sa, '안전')
    cw(t5.rows[2].cells[1],
       f'{sa[0]}위험성평가표   {sa[1]} 비상시 대응조직체계   {sa[2]}작업계획서 평가 및 관리계획')
    ca(t5.rows[2].cells[1], f'{sa[3]}기타 (                    )')
    tight(doc)

    # 하단 주석
    tn = mk(doc, 1, 1, [PW])
    cn = tn.rows[0].cells[0]
    cw(cn, NOTE_LINES[0], sz=8, bold=True)
    for line in NOTE_LINES[1:]:
        ca(cn, line, sz=8)
    tight(doc)

tight(doc)
out = r'C:\Users\user\Desktop\전기작업계획서.docx'
doc.save(out)
print('저장 완료:', out)
