# -*- coding: utf-8 -*-
"""작업안전 체크시트 엑셀을 웹에서 쓸 자료로 옮긴다.

엑셀의 2-x(관리) · 3-x(설비) · 4-x(작업) 시트를 읽어
sop-new/assets/checks.js 를 만든다. 사진은 넣지 않는다.

쓰는 법
    python sop_checksheet_import.py --xlsx "<엑셀 경로>"
    python sop_checksheet_import.py --xlsx "<...>" --only 4-1   한 장만 보기
"""
import argparse
import io
import json
import os
import re
import sys

HERE = os.path.dirname(os.path.abspath(__file__))
if HERE not in sys.path:
    sys.path.insert(0, HERE)
import sop_ppe_audit as A                                      # noqa: E402

ROOT = os.path.dirname(HERE)
OUT = os.path.join(ROOT, "sop-new", "assets", "checks.js")

묶음 = {"2": ("관리", "2. 관리 체크시트"),
       "3": ("설비", "3. 설비 체크시트"),
       "4": ("작업", "4. 작업 체크시트")}

# 체크박스 글에서 뽑은 이름 → 사이트에서 쓰는 보호구 이름
보호구이름 = {
    "안전화": "안전화", "안전모": "안전모", "안전벨트": "안전대",
    "안전대": "안전대", "보안경": "보안경", "보안면": "안면보호구",
    "보안경/보안면": "보안경", "미끄럼방지장화": "논슬립 장화",
    "방진마스크": "방진마스크", "방독마스크": "방독마스크",
    "귀마개": "귀마개", "귀덮개": "귀덮개", "절연장갑": "절연장갑",
    "절연화": "절연 안전화", "내열장갑": "내열장갑", "가죽장갑": "용접용 가죽장갑",
    "용접면": "용접용 보안면", "차광보안면": "용접용 보안면",
    "보호장갑": "보호장갑", "안전장갑": "안전장갑", "고무장갑": "고무장갑",
    "공기호흡기": "공기호흡기", "송기마스크": "송기마스크",
    "방열복": "보호의", "보호의": "보호의", "안전난간": None,
    "앞치마": "용접용 앞치마", "무릎보호대": "팔, 무릎보호대",
}


def 글(v):
    if v is None:
        return ""
    s = str(v).replace("∙", "·").replace("\x0b", " ")
    return re.sub(r"\s+", " ", s).strip()


def 한장(ws):
    rows = [[글(c) for c in r]
            for r in ws.iter_rows(max_row=60, max_col=9, values_only=True)]
    # 시트마다 내용이 A열에서 시작하기도, B열에서 시작하기도 한다.
    # 비어 있는 왼쪽 칸을 한꺼번에 떼어 기준을 맞춘다.
    while rows and all(not r[0] for r in rows):
        rows = [r[1:] + [""] for r in rows]

    제목 = ""
    for r in rows[:5]:
        for c in r:
            if len(c) > 8 and "CHECK" in c.upper():
                제목 = c
                break
        if 제목:
            break
    if not 제목:
        for r in rows[:5]:
            for c in r:
                if len(c) > 6:
                    제목 = c
                    break
            if 제목:
                break

    머리, 항목, 공통, 보호구, 기타 = [], [], [], [], ""
    공통시작 = False
    소제목 = ""
    for r in rows:
        a, b = r[0], r[1]
        나머지 = " ".join(x for x in r[1:] if x)

        if "□" in 나머지:
            for t in re.split(r"□", 나머지):
                t = t.strip(" ∙·")
                if not t:
                    continue
                t = re.sub(r"\(.*?\)", "", t).strip()
                이름 = 보호구이름.get(t.replace(" ", ""))
                if 이름 and 이름 not in 보호구:
                    보호구.append(이름)
            continue

        if a.startswith("공통사항"):
            공통시작 = True
            if re.match(r"^\d\)", b):
                공통.append(b)
            continue
        if a.startswith("기타사항"):
            공통시작 = False
            기타 = 나머지
            continue
        if 공통시작 and re.match(r"^\d\)", b):
            공통.append(b)
            continue

        if re.fullmatch(r"\d+", a) and b:
            것 = {"n": int(a), "text": b}
            if 소제목:
                것["sec"] = 소제목
                소제목 = ""
            항목.append(것)
            continue

        # 3-24 처럼 번호 없이 소제목만 있는 줄
        if a and not b and not 공통시작 and 항목 and len(a) < 30                 and a not in ("NO", "구분", "기타사항"):
            소제목 = a
            continue

        if a and b and not re.fullmatch(r"\d+", a) and len(a) < 14 \
                and a not in ("NO", "구분") and not 공통시작 and not 항목:
            머리.append([a, b])
            # 같은 줄 오른쪽에도 항목이 있으면 같이 담는다
            for i in (3, 4, 5):
                if i + 1 < len(r) and r[i] and len(r[i]) < 16 and r[i + 1]:
                    머리.append([r[i], r[i + 1]])
                    break

    return {"title": 제목, "head": 머리, "items": 항목,
            "common": 공통, "ppe": 보호구, "etc": 기타}


def 보호구뽑기(d):
    """엑셀 체크박스는 모든 장이 같은 공용 양식이라 쓸 수 없다.
    제목과 확인항목에 적힌 글에서 위험요인을 찾아 필요한 보호구를 끌어낸다."""
    if d.get("group") == "관리":
        return []                      # 서류 점검이라 보호구가 없다
    문장 = [d["title"]] + [x["text"] for x in d["items"]]
    이름글 = d["title"]
    필요 = ["안전모", "안전화"]         # 현장 점검의 기본
    for 이름, 낱말, 보호구 in A.위험:
        이름맞음 = [w for w in 낱말 if w in 이름글]
        본문맞음 = set()
        for 문 in 문장:
            if any(g in 문 for g in A.금지말):
                continue
            본문맞음 |= {w for w in 낱말 if w in 문}
        if 이름맞음 or len(본문맞음) >= 2:
            for x in 보호구:
                이름 = {"차광보안면": "용접용 보안면",
                      "가죽장갑": "용접용 가죽장갑"}.get(x, x)
                if 이름 not in 필요:
                    필요.append(이름)
    if "보호장갑" not in 필요 and not any("장갑" in x for x in 필요):
        필요.insert(2, "보호장갑")
    return 필요


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--xlsx", required=True)
    ap.add_argument("--only")
    a = ap.parse_args()

    from openpyxl import load_workbook
    wb = load_workbook(a.xlsx, read_only=True, data_only=True)

    # 목록 시트에서 이름을 가져온다
    이름표 = {}
    for n in wb.sheetnames:
        if not re.match(r"^(관리|설비|작업)리스트\d+$", n):
            continue
        for r in wb[n].iter_rows(max_col=3, values_only=True):
            k, v = 글(r[0]), 글(r[1]) if len(r) > 1 else ""
            if re.fullmatch(r"\d+-\d+", k) and v:
                이름표[k] = v

    out = {}
    for n in wb.sheetnames:
        키 = n
        if n == "109":
            키 = "2-9"
        if not re.fullmatch(r"\d+-\d+", 키):
            continue
        if a.only and 키 != a.only:
            continue
        d = 한장(wb[n])
        g = 묶음.get(키.split("-")[0])
        if not g:
            continue
        d["id"] = 키
        d["no"] = 키
        d["group"] = g[0]
        d["ppe"] = 보호구뽑기(d)
        d["name"] = 이름표.get(키) or re.sub(r"^\S+\s*", "", d["title"]) \
            .replace("CHECK SHEET", "").strip()
        out[키] = d

    if a.only:
        print(json.dumps(out, ensure_ascii=False, indent=1)[:2600])
        return

    글머리 = ("/* 안전 체크시트 자료 — 06-02.작업안전체크시트 에서 옮김. 사진 없음. */\n"
            "window.CHECKGROUPS = %s;\n\n"
            % json.dumps([{"id": k, "title": v[1], "name": v[0]}
                          for k, v in sorted(묶음.items())],
                         ensure_ascii=False, indent=1))
    io.open(OUT, "w", encoding="utf-8", newline="\n").write(
        글머리 + "window.CHECKS = " +
        json.dumps(out, ensure_ascii=False, indent=1) + ";\n")

    print("체크시트 %d장 저장 : %s (%.0f KB)"
          % (len(out), OUT, os.path.getsize(OUT) / 1024))
    for g, (nm, _t) in sorted(묶음.items()):
        아이들 = [k for k in out if k.startswith(g + "-")]
        항목수 = sum(len(out[k]["items"]) for k in 아이들)
        보호구수 = sum(len(out[k]["ppe"]) for k in 아이들)
        print("  %-3s %2d장 / 확인항목 %3d개 / 보호구 %3d건"
              % (nm, len(아이들), 항목수, 보호구수))


if __name__ == "__main__":
    main()
