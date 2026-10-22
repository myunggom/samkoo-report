# -*- coding: utf-8 -*-
"""하자 치유 리스트 — 전체본의 완료 표시를 통합본으로 옮긴다.

규칙
    통합본 한 행의 R열(통합번호)에 적힌 전체본 번호가 **전부** 완료(O)여야
    그 행의 O열(확인)에 O 를 넣는다. 하나라도 미완료면 비운다.
    R열이 비어 있는 행은 묶이지 않은 단일 건이므로 자기 B열 번호로 대조한다.

왜 엑셀을 직접 띄우는가
    두 파일에 치유사진이 수천 장 박혀 있다. 파이썬 라이브러리로 저장하면
    사진과 도면이 전부 날아간다. 실제 엑셀로 바꿀 칸만 고친다.

원본은 고치지 않는다
    항상 새 이름으로 사본을 만들어 사본에만 쓴다.

쓰는 법
    python haja_merge_sync.py --전체 "<전체본.xlsx>" --통합 "<통합본.xlsx>"
    python haja_merge_sync.py ... --대조          쓰지 않고 바뀔 것만 보여준다
    python haja_merge_sync.py ... --유지          이미 있는 O 는 지우지 않는다
"""
import argparse
import datetime as dt
import os
import re
import shutil
import sys

머리 = 6                      # 6행이 머리글, 자료는 7행부터
C공종, C번호, C확인, C통합번호 = 1, 2, 15, 18       # A, B, O, R


def 텍스트(v):
    return "" if v is None else str(v).strip()


def 정수(v):
    try:
        return int(float(텍스트(v)))
    except (TypeError, ValueError):
        return None


def 시트짝(전체시트, 통합시트):
    """'건축' ↔ '건축_통합' 처럼 이름으로 짝을 맞춘다."""
    짝 = []
    남은 = list(통합시트)
    for s in 전체시트:
        for t in list(남은):
            if t == s or t == s + "_통합" or t.replace("_통합", "") == s:
                짝.append((s, t))
                남은.remove(t)
                break
    return 짝


def 전체읽기(path):
    """전체본 → {시트: {(공종, 번호): 확인값}}"""
    from openpyxl import load_workbook
    wb = load_workbook(path, read_only=True, data_only=True)
    out = {}
    try:
        for s in wb.sheetnames:
            표 = {}
            for row in wb[s].iter_rows(min_row=머리 + 1, max_col=20,
                                       values_only=True):
                n = 정수(row[C번호 - 1])
                if n is None:
                    continue
                표[(텍스트(row[C공종 - 1]).upper(), n)] = 텍스트(row[C확인 - 1])
            if 표:
                out[s] = 표
    finally:
        wb.close()
    return out


def 통합읽기(path):
    """통합본 → {시트: [{행, 공종, 번호, 확인, 통합번호}, ...]}"""
    from openpyxl import load_workbook
    wb = load_workbook(path, read_only=True, data_only=True)
    out = {}
    try:
        for s in wb.sheetnames:
            행들 = []
            for i, row in enumerate(wb[s].iter_rows(min_row=머리 + 1,
                                                    max_col=20,
                                                    values_only=True),
                                    머리 + 1):
                번호 = 정수(row[C번호 - 1])
                통합 = 텍스트(row[C통합번호 - 1])
                if 번호 is None and not 통합:
                    continue
                행들.append({"행": i, "공종": 텍스트(row[C공종 - 1]).upper(),
                            "번호": 번호, "확인": 텍스트(row[C확인 - 1]),
                            "통합번호": 통합})
            if 행들:
                out[s] = 행들
        return out
    finally:
        wb.close()


def 번호쪼개기(글):
    """'A154, A155' → [('A',154), ('A',155)]. 해석 못한 조각도 돌려준다."""
    좋음, 나쁨 = [], []
    for t in re.split(r"[,\s/·]+", 글):
        t = t.strip()
        if not t:
            continue
        m = re.fullmatch(r"([A-Za-z]+)\s*(\d+)", t)
        if m:
            좋음.append((m.group(1).upper(), int(m.group(2))))
        else:
            나쁨.append(t)
    return 좋음, 나쁨


def 계산(전체, 통합, 유지=False):
    """시트별로 바꿀 칸과 경고를 돌려준다."""
    결과, 경고 = {}, []
    for s1, s2 in 시트짝(list(전체), list(통합)):
        표 = 전체[s1]
        바꿀것, 쓴키 = [], set()
        통계 = {"새로O": 0, "O지움": 0, "그대로": 0}
        for r in 통합[s2]:
            쪽, 나쁨 = ([], [])
            if r["통합번호"]:
                쪽, 나쁨 = 번호쪼개기(r["통합번호"])
                for x in 나쁨:
                    경고.append("%s %d행 통합번호 해석 불가 : %s"
                                % (s2, r["행"], x))
            elif r["번호"] is not None:
                쪽 = [(r["공종"], r["번호"])]
            if not 쪽:
                경고.append("%s %d행 : 대조할 번호가 없어 건드리지 않음"
                            % (s2, r["행"]))
                continue

            값 = []
            for g, n in 쪽:
                쓴키.add((g, n))
                if (g, n) not in 표:
                    경고.append("%s %d행 : 전체본에 %s%d 가 없음"
                                % (s2, r["행"], g, n))
                    값.append(None)
                else:
                    값.append(표[(g, n)].upper() == "O")
            다됨 = bool(값) and all(v is True for v in 값)
            지금 = r["확인"].upper() == "O"
            if 다됨 == 지금:
                통계["그대로"] += 1
                continue
            if 다됨:
                통계["새로O"] += 1
                바꿀것.append((r["행"], "O", r, 쪽, 값))
            else:
                if 유지:
                    통계["그대로"] += 1
                    continue
                통계["O지움"] += 1
                바꿀것.append((r["행"], None, r, 쪽, 값))

        안쓴키 = [k for k in 표 if k not in 쓴키]
        if 안쓴키:
            보기 = ", ".join("%s%d" % k for k in sorted(안쓴키,
                                                      key=lambda x: x[1])[:8])
            경고.append("%s : 통합본이 가리키지 않는 전체본 번호 %d건 (%s …)"
                        % (s2, len(안쓴키), 보기))
        결과[s2] = {"바꿀것": 바꿀것, "통계": 통계}
    return 결과, 경고


def 쓰기(사본, 결과):
    """엑셀을 띄워 바꿀 칸만 고친다. 집계·완료율은 수식이라 알아서 다시 센다."""
    import xlwings as xw

    app = xw.App(visible=False, add_book=False)
    app.display_alerts = False
    app.screen_updating = False
    try:
        wb = app.books.open(os.path.abspath(사본))
        try:
            for s, d in 결과.items():
                if not d["바꿀것"]:
                    continue
                ws = wb.sheets[s]
                for 행, 값, _r, _쪽, _값들 in d["바꿀것"]:
                    ws.range((행, C확인)).value = 값
            app.calculate()
            wb.save()
        finally:
            wb.close()
    finally:
        app.quit()


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--전체", required=True, help="전체 번호가 있는 파일")
    ap.add_argument("--통합", required=True, help="번호 통합 파일")
    ap.add_argument("--출력", help="만들 사본 (없으면 통합 파일명의 날짜를 오늘로)")
    ap.add_argument("--유지", action="store_true",
                    help="이미 있는 O 는 지우지 않는다")
    ap.add_argument("--대조", action="store_true", help="쓰지 않고 보여만 준다")
    a = ap.parse_args()

    print("전체본 : %s" % os.path.basename(a.전체))
    print("통합본 : %s" % os.path.basename(a.통합))
    print()

    전체 = 전체읽기(a.전체)
    통합 = 통합읽기(a.통합)
    결과, 경고 = 계산(전체, 통합, 유지=a.유지)

    합 = {"새로O": 0, "O지움": 0, "그대로": 0}
    for s, d in 결과.items():
        t = d["통계"]
        for k in 합:
            합[k] += t[k]
        print("## %-12s 새로 O %3d / O 지움 %3d / 그대로 %4d"
              % (s, t["새로O"], t["O지움"], t["그대로"]))
        for 행, 값, r, 쪽, 값들 in d["바꿀것"][:4]:
            미완 = ["%s%d" % k for k, v in zip(쪽, 값들) if v is not True]
            print("     %4d행 %-4s %s  (%d건 중 완료 %d건%s)"
                  % (행, "→ O" if 값 else "→ 지움",
                     (r["통합번호"] or "단일")[:40], len(쪽),
                     sum(1 for v in 값들 if v),
                     ", 미완료 " + ", ".join(미완[:4]) if 미완 else ""))
        if len(d["바꿀것"]) > 4:
            print("     ... 그 외 %d행" % (len(d["바꿀것"]) - 4))
    print("\n합계 : 새로 O %d / O 지움 %d / 그대로 %d"
          % (합["새로O"], 합["O지움"], 합["그대로"]))

    if 경고:
        print("\n확인할 것 %d건" % len(경고))
        for x in 경고[:12]:
            print("   ! %s" % x)
        if len(경고) > 12:
            print("   ... 그 외 %d건" % (len(경고) - 12))

    if a.대조:
        print("\n(대조만 했습니다. 아무것도 쓰지 않았습니다.)")
        return

    out = a.출력
    if not out:
        오늘 = dt.date.today().strftime("%Y.%m.%d")
        이름 = re.sub(r"\d{4}\.\d{2}\.\d{2}", 오늘,
                    os.path.basename(a.통합))
        if 오늘 not in 이름:
            뿌리, 꼬리 = os.path.splitext(os.path.basename(a.통합))
            이름 = "%s_%s%s" % (뿌리, 오늘, 꼬리)
        out = os.path.join(os.path.dirname(a.통합), 이름)

    if os.path.abspath(out) == os.path.abspath(a.통합):
        sys.exit("사본 이름이 원본과 같습니다. --출력 으로 다른 이름을 주세요.")
    print("\n사본 만드는 중 ... %s" % os.path.basename(out))
    shutil.copy2(a.통합, out)
    쓰기(out, 결과)
    print("저장 완료 : %s" % out)
    print("※ 원본 두 개는 읽기만 했습니다.")


if __name__ == "__main__":
    main()
