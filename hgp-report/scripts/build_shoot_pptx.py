# -*- coding: utf-8 -*-
"""
촬영 완료보고서 PPT 양식 — 웹 PDF(components/shoot/ReportDocument.tsx)와 "똑같은" 모양으로 만든다.

좌표·크기·색·글자 크기는 PDF 페이지(1123×794px, A4 가로)를 브라우저에서 실측한 값을 그대로 쓴다.
PDF 디자인을 바꾸면 이 파일의 좌표도 같이 맞춰야 한다.

  1  표지        건물 사진 + 네이비 그라디언트 + Samkoo 로고 + 제목 블록 + 하단 바
  2  촬영 개요   개요 표(4행) + 촬영 회차 표(이번 회차 강조)
  3  구역 사진   정보표 + 구역 배지 + 사진 4×2 + 촬영 중 특이사항·비고   ← 웹에서 구역·8장마다 복제
  4  특이사항    정보표 + 사진 2×2 + 특이사항 설명

출력
  python scripts/build_shoot_pptx.py            → public/templates/shoot.pptx (웹 출력용, {태그} 포함)
  python scripts/build_shoot_pptx.py blank OUT  → OUT (직접 쓰는 빈 양식, 태그 없음, 구역 3장)

웹 채우기: lib/shoot/pptx.ts — 사진칸 크기(SECTION_BOX/SPECIAL_BOX)와 배지 이름(BADGE)이 여기와 맞아야 함.
"""
import os
import sys

from lxml import etree as ET
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Pt

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
COVER_IMG = os.path.join(ROOT, "public", "cover-building.jpg")
LOGO = os.path.join(ROOT, "public", "samkoo.png")
OUT = os.path.join(ROOT, "public", "templates", "shoot.pptx")

FONT = "맑은 고딕"
PAGE_W, PAGE_H = 1123, 794
NAVY = "0F2C5C"
SECTIONS = ["보양 및 세팅", "촬영", "철수 및 정리"]
MODE = "template"  # or "blank"


def px(v):
    return Emu(int(round(v * 9525)))


def fpt(px_size):
    """CSS px → pt"""
    return Pt(px_size * 0.75)


def T(tag, blank=""):
    return "{%s}" % tag if MODE == "template" else blank


def rgb(h):
    return RGBColor.from_string(h)


# ── 글자 ────────────────────────────────────────────────────
def style_run(r, size_px, color="111111", bold=False, spacing_px=0):
    f = r.font
    f.size = fpt(size_px)
    f.bold = bold
    f.name = FONT
    f.color.rgb = rgb(color)
    rpr = r._r.get_or_add_rPr()
    for tag in ("a:ea", "a:cs"):
        e = rpr.find(qn(tag))
        if e is None:
            e = ET.SubElement(rpr, qn(tag))
        e.set("typeface", FONT)
    if spacing_px:
        rpr.set("spc", str(int(round(spacing_px * 0.75 * 100))))


def R(s, size=13, color="111111", bold=False, sp=0):
    return (s, size, color, bold, sp)


def text(slide, x, y, w, h, runs, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.MIDDLE, pad=(0, 0, 0, 0), line=None, name=None):
    """runs: [R(...)] 한 문단."""
    tb = slide.shapes.add_textbox(px(x), px(y), px(w), px(h))
    if name:
        tb.name = name
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left, tf.margin_top, tf.margin_right, tf.margin_bottom = [px(v) for v in pad]
    tf.auto_size = None
    p = tf.paragraphs[0]
    p.alignment = align
    if line:
        p.line_spacing = line
    for s, size, color, bold, sp in runs:
        r = p.add_run()
        r.text = s
        style_run(r, size, color, bold, sp)
    return tb


# ── 도형 ────────────────────────────────────────────────────
def shape(slide, x, y, w, h, fill=None, line=None, radius=0, alpha=None, name=None):
    kind = MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE
    s = slide.shapes.add_shape(kind, px(x), px(y), px(w), px(h))
    if name:
        s.name = name
    if radius:
        s.adjustments[0] = min(0.5, radius / min(w, h))
    if fill:
        s.fill.solid()
        s.fill.fore_color.rgb = rgb(fill)
        if alpha is not None:
            clr = s._element.spPr.find(qn("a:solidFill")).find(qn("a:srgbClr"))
            ET.SubElement(clr, qn("a:alpha"), val=str(int(alpha * 100000)))
    else:
        s.fill.background()
    if line:
        s.line.color.rgb = rgb(line)
        s.line.width = px(1)
    else:
        s.line.fill.background()
    s.shadow.inherit = False
    tf = s.text_frame
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    return s


def gradient(s, stops, angle_deg):
    """stops: [(pos 0~1, hex, alpha 0~1)]. angle: DrawingML(0=왼→오, 시계방향)."""
    spPr = s._element.spPr
    for tag in ("a:solidFill", "a:noFill", "a:gradFill"):
        e = spPr.find(qn(tag))
        if e is not None:
            spPr.remove(e)
    g = ET.Element(qn("a:gradFill"), rotWithShape="1")
    gs = ET.SubElement(g, qn("a:gsLst"))
    for pos, hx, a in stops:
        st = ET.SubElement(gs, qn("a:gs"), pos=str(int(pos * 100000)))
        c = ET.SubElement(st, qn("a:srgbClr"), val=hx)
        ET.SubElement(c, qn("a:alpha"), val=str(int(a * 100000)))
    ET.SubElement(g, qn("a:lin"), ang=str(int(angle_deg * 60000)), scaled="0")
    spPr.find(qn("a:prstGeom")).addnext(g)


def picture_cover(slide, path, x, y, w, h):
    """object-fit: cover 처럼 잘라 넣기"""
    from PIL import Image

    iw, ih = Image.open(path).size
    pic = slide.shapes.add_picture(path, px(x), px(y), px(w), px(h))
    box, img = w / h, iw / ih
    if img > box:
        c = (1 - box / img) / 2
        pic.crop_left = pic.crop_right = c
    else:
        c = (1 - img / box) / 2
        pic.crop_top = pic.crop_bottom = c
    return pic


# ── 표 ─────────────────────────────────────────────────────
def cell_style(cell, runs, fill="FFFFFF", align=PP_ALIGN.LEFT, pad=(12, 7)):
    cell.text = ""
    tf = cell.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = align
    for s, size, color, bold, sp in runs:
        r = p.add_run()
        r.text = s or " "  # 빈 칸도 글자 크기를 유지해야 행 높이가 PDF와 같음
        style_run(r, size, color, bold, sp)
    cell.vertical_anchor = MSO_ANCHOR.MIDDLE
    cell.margin_left = cell.margin_right = px(pad[0])
    cell.margin_top = cell.margin_bottom = px(pad[1])
    cell.fill.solid()
    cell.fill.fore_color.rgb = rgb(fill)
    tcPr = cell._tc.get_or_add_tcPr()
    for tag in ("a:lnL", "a:lnR", "a:lnT", "a:lnB"):
        old = tcPr.find(qn(tag))
        if old is not None:
            tcPr.remove(old)
        ln = ET.SubElement(tcPr, qn(tag), w="9525", cap="flat", cmpd="sng", algn="ctr")
        sf = ET.SubElement(ln, qn("a:solidFill"))
        ET.SubElement(sf, qn("a:srgbClr"), val="CBD5E1")
        ET.SubElement(ln, qn("a:prstDash"), val="solid")
    f = tcPr.find(qn("a:solidFill"))  # 스키마 순서: 테두리 다음 채우기
    tcPr.remove(f)
    tcPr.append(f)


def table(slide, x, y, widths, heights):
    shp = slide.shapes.add_table(len(heights), len(widths), px(x), px(y), px(sum(widths)), px(sum(heights)))
    tbl = shp.table
    tblPr = tbl._tbl.tblPr
    for a in ("firstRow", "bandRow", "firstCol", "lastRow", "lastCol", "bandCol"):
        tblPr.set(a, "0")
    sid = tblPr.find(qn("a:tableStyleId"))
    if sid is not None:
        tblPr.remove(sid)
    for i, w in enumerate(widths):
        tbl.columns[i].width = px(w)
    for i, h in enumerate(heights):
        tbl.rows[i].height = px(h)
    return tbl


TH = dict(fill="EEF2F7", align=PP_ALIGN.CENTER)


def th(s, size=13):
    return [R(s, size, "334155", True)]


# ── 공통 부분 ───────────────────────────────────────────────
def page_header(slide, title, subtitle_runs):
    text(slide, 28, 28, 700, 37.5, [R(title, 25, NAVY, True)])
    text(slide, 28, 68.5, 900, 18, subtitle_runs)
    slide.shapes.add_picture(LOGO, px(1013.1), px(28), px(81.9), px(38))


def info_compact(slide):
    tw, dw = 92, (1067 - 92 * 3) / 3
    tbl = table(slide, 28, 98.5, [tw, dw, tw, dw, tw, dw], [34.5, 34.5])
    rows = [
        [("촬영종류", 1), (T("촬영종류"), 0), ("촬영명", 1), (T("촬영명"), 0), ("제작사", 1), (T("제작사"), 0)],
        [("관리자", 1), (T("관리자"), 0), ("보양 및 세팅", 1), (T("보양세팅"), 0), ("촬영 및 철수", 1), (T("촬영철수"), 0)],
    ]
    for ri, row in enumerate(rows):
        for ci, (s, is_th) in enumerate(row):
            if is_th:
                # PPT 글꼴(맑은 고딕)이 PDF 글꼴보다 넓어 좌우 여백을 줄여야 한 줄에 들어감
                cell_style(tbl.cell(ri, ci), th(s), pad=(4, 7), **TH)
            else:
                cell_style(tbl.cell(ri, ci), [R(s, 13)])


def photo_cell(slide, x, y, w, h, n, tag):
    """PDF PhotoCell: 둥근 틀(#f8fafc) + '사진 없음' + 사진(캡션은 사진에 합성) + 번호 배지"""
    shape(slide, x, y, w, h, fill="F8FAFC", line="E2E8F0", radius=6)
    text(slide, x + 1, y + 1, w - 2, h - 2, [R("사진 없음", 12, "CBD5E1")], align=PP_ALIGN.CENTER)
    if MODE == "template":
        # docxtemplater 이미지 모듈이 이 글상자 자리에 사진을 넣음(크기는 lib/shoot/pptx.ts)
        text(slide, x + 1, y + 1, w - 2, h - 2, [R("{%%%s}" % tag, 8, "F8FAFC")], anchor=MSO_ANCHOR.TOP)
    b = shape(slide, x + 6, y + 6, 18, 18, fill=NAVY, radius=4, alpha=0.9)
    p = b.text_frame.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    b.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    r = p.add_run()
    r.text = str(n)
    style_run(r, 11, "FFFFFF", True)


def note_runs(tag, cond):
    """값 있으면 검정, 없으면 회색 '해당 없음' (PDF와 동일)"""
    if MODE == "template":
        return [R("{#%s}{%s}{/%s}" % (cond, tag, cond), 13, "111111"), R("{^%s}해당 없음{/%s}" % (cond, cond), 13, "94A3B8")]
    return [R("", 13)]


def blank_slide(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    for ph in list(s.placeholders):
        ph._element.getparent().remove(ph._element)
    return s


# ── 슬라이드 ────────────────────────────────────────────────
def build_cover(prs):
    s = blank_slide(prs)
    picture_cover(s, COVER_IMG, 0, 0, PAGE_W, PAGE_H)
    if MODE == "template":
        # 보고서에 겉표지 사진을 따로 올렸으면 그 사진으로 덮음
        text(s, 0, 0, PAGE_W, PAGE_H, [R("{%cover}", 8, "FFFFFF")], anchor=MSO_ANCHOR.TOP)
    g = shape(s, 0, 0, PAGE_W, PAGE_H, fill="091833")
    # CSS linear-gradient(100deg, …) = DrawingML 10°
    gradient(g, [(0, "091833", 0.94), (0.34, "0C2148", 0.82), (0.60, NAVY, 0.35), (0.82, NAVY, 0.0), (1, NAVY, 0.0)], 10)
    s.shapes.add_picture(LOGO, px(979.8), px(36), px(99.2), px(46))
    text(s, 60, 238.2, 640, 24, [R("한독 · 제넥신 · 프로젠 연구소", 16, "CFE0F5", True, 3)])
    text(s, 60, 266.2, 640, 22.5, [R(T("표지종류", "촬영"), 15, "9DB8DD", False, 2)])
    text(s, 60, 302.7, 900, 63.8, [R("촬영 완료보고서", 58, "FFFFFF", True)])
    text(s, 60, 372.5, 900, 45, [R(T("촬영명"), 30, "EAF1FB", True)])
    shape(s, 60, 441.5, 90, 4, fill="4F86D6", radius=2)
    text(s, 60, 469.5, 900, 28.8, [R("제작사 : " + T("제작사"), 16, "DBE6F6")])
    text(s, 60, 498.3, 900, 28.8, [R("촬영 기간 : " + T("촬영기간"), 16, "DBE6F6")])
    manager = "{#관리자있음}관리자 : {관리자}{/관리자있음}" if MODE == "template" else "관리자 : "
    text(s, 60, 527, 900, 28.8, [R(manager, 16, "DBE6F6")])
    bar = shape(s, 0, PAGE_H - 8, PAGE_W, 8, fill=NAVY)
    gradient(bar, [(0, NAVY, 1), (1, "4F86D6", 1)], 0)


def build_overview(prs):
    s = blank_slide(prs)
    page_header(s, "촬영 개요", [R("촬영 완료보고서 · 개요 안내", 12, "64748B")])
    tbl = table(s, 28, 98.5, [150, 383.5, 150, 383.5], [51.5] * 4)
    rows = [
        ("촬영종류", "촬영종류", "촬영명", "촬영명"),
        ("제작사", "제작사", "관리자", "관리자"),
        ("촬영 일시", "촬영일시", "촬영 기간", "촬영기간"),
        ("보양 및 세팅", "보양세팅", "촬영 및 철수", "촬영철수"),
    ]
    for ri, (a, ta, b, tb) in enumerate(rows):
        cell_style(tbl.cell(ri, 0), th(a, 15), pad=(14, 14), **TH)
        cell_style(tbl.cell(ri, 1), [R(T(ta), 15)], pad=(14, 14))
        cell_style(tbl.cell(ri, 2), th(b, 15), pad=(14, 14), **TH)
        cell_style(tbl.cell(ri, 3), [R(T(tb), 15)], pad=(14, 14))

    shape(s, 28, 335.3, 5, 18, fill=NAVY, radius=2)
    text(s, 41, 331.5, 300, 25.5, [R("촬영 회차", 17, NAVY, True)])
    if MODE == "template":
        summary = [R("총 {회차수}회차 · 이번은 ", 13, "64748B"), R("{이번}번째", 13, NAVY, True), R(" 촬영", 13, "64748B")]
    else:
        summary = [R("총     회차 · 이번은     번째 촬영", 13, "64748B")]
    text(s, 595, 338.3, 500, 19.5, summary, align=PP_ALIGN.RIGHT)

    widths = [80, 477, 260, 130, 120]
    n_body = 2 if MODE == "template" else 5
    tb = table(s, 28, 367.8, widths, [34.5] * (1 + n_body))
    for ci, h in enumerate(["회차", "촬영 일시", "촬영 및 철수", "관리자", "비고"]):
        cell_style(tb.cell(0, ci), th(h), **TH)
    aligns = [PP_ALIGN.CENTER, PP_ALIGN.LEFT, PP_ALIGN.LEFT, PP_ALIGN.CENTER, PP_ALIGN.CENTER]
    if MODE == "template":
        # 두 행을 한 묶음으로 반복: 이번 회차면 강조 행({#now}), 아니면 보통 행({^now})
        now_vals = ["{#회차}{#now}{회차명}", "{일시}", "{철수}", "{관리자}", "이번 촬영{/now}"]
        other_vals = ["{^now}{회차명}", "{일시}", "{철수}", "{관리자}", "{/now}{/회차}"]
        for ci, v in enumerate(now_vals):
            color = NAVY if ci in (0, 4) else "111111"
            cell_style(tb.cell(1, ci), [R(v, 13 if ci != 4 else 11, color, True)], fill="E8F0FB", align=aligns[ci])
        for ci, v in enumerate(other_vals):
            cell_style(tb.cell(2, ci), [R(v, 13, NAVY if ci == 0 else "111111", ci == 0)], align=aligns[ci])
    else:
        for ri in range(1, 1 + n_body):
            for ci in range(5):
                cell_style(tb.cell(ri, ci), [R("", 13)], align=aligns[ci])


def build_section(prs, section_name=None):
    s = blank_slide(prs)
    page_header(s, "촬영 완료보고서", [R("촬영 기간 : " + T("촬영기간"), 12, "64748B")])
    info_compact(s)
    # 구역 배지 — 폭은 글자 길이에 맞춰 lib/shoot/pptx.ts가 조정(BADGE), (1/2)도 그 옆으로(BADGE_PART)
    b = shape(s, 28, 180.5, 100.3, 29, fill=NAVY, radius=4, name="BADGE")
    p = b.text_frame.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    b.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    r = p.add_run()
    r.text = T("구역", section_name or "")
    style_run(r, 14, "FFFFFF", True)
    text(s, 136.3, 186, 100, 18, [R(T("쪽"), 12, "94A3B8")], name="BADGE_PART")
    for i in range(8):
        cx = 28 + (i % 4) * 268.77
        cy = 217.5 + (i // 4) * 221.5
        photo_cell(s, cx, cy, 260.8, 213.5, i + 1, "p%d" % (i + 1))
    shape(s, 28, 662.5, 1067, 103.5, line="E2E8F0", radius=8)
    shape(s, 29, 663.5, 1065, 31.5, fill="EEF2F7")
    text(s, 29, 663.5, 1065, 31.5, [R("촬영 중 특이사항 · 비고", 13, "334155", True)], pad=(12, 0, 12, 0))
    text(s, 29, 695, 1065, 70, note_runs("비고", "noteOk"), anchor=MSO_ANCHOR.TOP, pad=(12, 10, 12, 10), line=1.25)


def build_special(prs):
    s = blank_slide(prs)
    page_header(s, "특이사항", [R("촬영 기간 : " + T("촬영기간"), 12, "64748B")])
    info_compact(s)
    for i in range(4):
        cx = 28 + (i % 2) * 346.4
        cy = 180.5 + (i // 2) * 297.8
        photo_cell(s, cx, cy, 336.4, 287.8, i + 1, "s%d" % (i + 1))
    shape(s, 722.9, 180.5, 372.1, 585.5, line="E2E8F0", radius=8)
    shape(s, 723.9, 181.5, 370.1, 38, fill="EEF2F7")
    text(s, 723.9, 181.5, 370.1, 38, [R("특이사항 설명", 14, "334155", True)], pad=(12, 0, 12, 0))
    text(s, 723.9, 219.5, 370.1, 545.5, note_runs("특이설명", "specialOk"), anchor=MSO_ANCHOR.TOP, pad=(12, 12, 12, 12), line=1.35)


def build(out_path):
    prs = Presentation()
    prs.slide_width, prs.slide_height = px(PAGE_W), px(PAGE_H)
    build_cover(prs)
    build_overview(prs)
    if MODE == "template":
        build_section(prs)
    else:
        for name in SECTIONS:
            build_section(prs, name)
    build_special(prs)
    prs.save(out_path)
    print("saved", out_path, os.path.getsize(out_path))


if __name__ == "__main__":
    if len(sys.argv) >= 3 and sys.argv[1] == "blank":
        MODE = "blank"
        build(sys.argv[2])
    else:
        build(OUT)
